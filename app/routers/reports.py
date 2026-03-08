from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from datetime import date, timedelta
from decimal import Decimal
from typing import Optional
import io
import json

from app.database import get_db
from app.models import DailySale, DailyExpense, User
from app.routers.auth import get_current_user, require_owner

router = APIRouter(prefix="/api/reports", tags=["Reports"])


def to_float(val) -> float:
    if val is None:
        return 0.0
    return float(val)


@router.get("/profit-loss")
def profit_loss_report(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    today = date.today()
    start = date.fromisoformat(date_from) if date_from else today.replace(day=1)
    end = date.fromisoformat(date_to) if date_to else today

    sales_q = db.query(DailySale).filter(DailySale.date >= start, DailySale.date <= end)
    expenses_q = db.query(DailyExpense).filter(DailyExpense.date >= start, DailyExpense.date <= end)

    if current_user.role == "employee":
        sales_q = sales_q.filter(DailySale.employee_id == current_user.id)
        expenses_q = expenses_q.filter(DailyExpense.employee_id == current_user.id)

    total_sales = to_float(sales_q.with_entities(func.sum(DailySale.amount)).scalar())
    total_expenses = to_float(expenses_q.with_entities(func.sum(DailyExpense.amount)).scalar())

    # Sales by type
    sales_by_type = (
        sales_q.with_entities(DailySale.sale_type, func.sum(DailySale.amount), func.count())
        .group_by(DailySale.sale_type).all()
    )

    # Expenses by category
    expenses_by_category = (
        expenses_q.with_entities(DailyExpense.category, func.sum(DailyExpense.amount), func.count())
        .group_by(DailyExpense.category)
        .order_by(func.sum(DailyExpense.amount).desc()).all()
    )

    return {
        "period": {"from": start.isoformat(), "to": end.isoformat()},
        "total_sales": round(total_sales, 3),
        "total_expenses": round(total_expenses, 3),
        "net_profit": round(total_sales - total_expenses, 3),
        "sales_by_type": [
            {"type": t, "amount": round(to_float(a), 3), "count": c}
            for t, a, c in sales_by_type
        ],
        "expenses_by_category": [
            {"category": cat, "amount": round(to_float(a), 3), "count": c}
            for cat, a, c in expenses_by_category
        ],
    }


@router.get("/daily")
def daily_report(
    report_date: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    d = date.fromisoformat(report_date) if report_date else date.today()

    sales_q = db.query(DailySale).filter(DailySale.date == d)
    expenses_q = db.query(DailyExpense).filter(DailyExpense.date == d)

    if current_user.role == "employee":
        sales_q = sales_q.filter(DailySale.employee_id == current_user.id)
        expenses_q = expenses_q.filter(DailyExpense.employee_id == current_user.id)

    sales = sales_q.order_by(DailySale.id).all()
    expenses = expenses_q.order_by(DailyExpense.id).all()

    total_sales = sum(float(s.amount) for s in sales)
    total_expenses = sum(float(e.amount) for e in expenses)

    return {
        "date": d.isoformat(),
        "sales": [
            {
                "id": s.id, "customer_name": s.customer_name,
                "sale_type": s.sale_type, "description": s.description,
                "amount": str(s.amount), "payment_method": s.payment_method,
                "employee": s.employee.full_name if s.employee else None,
            }
            for s in sales
        ],
        "expenses": [
            {
                "id": e.id, "category": e.category,
                "description": e.description, "amount": str(e.amount),
                "payment_method": e.payment_method,
                "employee": e.employee.full_name if e.employee else None,
            }
            for e in expenses
        ],
        "total_sales": round(total_sales, 3),
        "total_expenses": round(total_expenses, 3),
        "net_profit": round(total_sales - total_expenses, 3),
    }


@router.get("/employee-activity")
def employee_activity_report(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    current_user: User = Depends(require_owner),
    db: Session = Depends(get_db),
):
    today = date.today()
    start = date.fromisoformat(date_from) if date_from else today.replace(day=1)
    end = date.fromisoformat(date_to) if date_to else today

    employees = db.query(User).filter(User.is_active == True).all()
    result = []
    for emp in employees:
        sales_count = (
            db.query(func.count(DailySale.id))
            .filter(DailySale.employee_id == emp.id, DailySale.date >= start, DailySale.date <= end)
            .scalar()
        )
        sales_total = to_float(
            db.query(func.sum(DailySale.amount))
            .filter(DailySale.employee_id == emp.id, DailySale.date >= start, DailySale.date <= end)
            .scalar()
        )
        expenses_count = (
            db.query(func.count(DailyExpense.id))
            .filter(DailyExpense.employee_id == emp.id, DailyExpense.date >= start, DailyExpense.date <= end)
            .scalar()
        )
        expenses_total = to_float(
            db.query(func.sum(DailyExpense.amount))
            .filter(DailyExpense.employee_id == emp.id, DailyExpense.date >= start, DailyExpense.date <= end)
            .scalar()
        )
        result.append({
            "employee_id": emp.id,
            "employee_name": emp.full_name,
            "role": emp.role,
            "sales_count": sales_count,
            "sales_total": round(sales_total, 3),
            "expenses_count": expenses_count,
            "expenses_total": round(expenses_total, 3),
        })

    return {
        "period": {"from": start.isoformat(), "to": end.isoformat()},
        "employees": result,
    }


@router.get("/cash-flow")
def cash_flow_report(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    today = date.today()
    start = date.fromisoformat(date_from) if date_from else today.replace(day=1)
    end = date.fromisoformat(date_to) if date_to else today

    sales_q = db.query(DailySale).filter(DailySale.date >= start, DailySale.date <= end)
    expenses_q = db.query(DailyExpense).filter(DailyExpense.date >= start, DailyExpense.date <= end)

    if current_user.role == "employee":
        sales_q = sales_q.filter(DailySale.employee_id == current_user.id)
        expenses_q = expenses_q.filter(DailyExpense.employee_id == current_user.id)

    # Sales by payment method
    sales_by_method = (
        sales_q.with_entities(DailySale.payment_method, func.sum(DailySale.amount))
        .group_by(DailySale.payment_method).all()
    )
    # Expenses by payment method
    expenses_by_method = (
        expenses_q.with_entities(DailyExpense.payment_method, func.sum(DailyExpense.amount))
        .group_by(DailyExpense.payment_method).all()
    )

    total_sales = to_float(sales_q.with_entities(func.sum(DailySale.amount)).scalar())
    total_expenses = to_float(expenses_q.with_entities(func.sum(DailyExpense.amount)).scalar())

    return {
        "period": {"from": start.isoformat(), "to": end.isoformat()},
        "total_income": round(total_sales, 3),
        "total_expenses": round(total_expenses, 3),
        "net_cash_flow": round(total_sales - total_expenses, 3),
        "income_by_method": [
            {"method": m, "amount": round(to_float(a), 3)} for m, a in sales_by_method
        ],
        "expenses_by_method": [
            {"method": m, "amount": round(to_float(a), 3)} for m, a in expenses_by_method
        ],
    }


@router.get("/export/excel")
def export_excel(
    report_type: str = "profit-loss",
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed. Run: pip install openpyxl")

    today = date.today()
    start = date.fromisoformat(date_from) if date_from else today.replace(day=1)
    end = date.fromisoformat(date_to) if date_to else today

    wb = openpyxl.Workbook()
    ws = wb.active

    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="714B67", end_color="714B67", fill_type="solid")
    title_font = Font(bold=True, size=14)
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    if report_type == "sales":
        ws.title = "Sales Report"
        ws.merge_cells('A1:G1')
        ws['A1'] = f"Sales Report ({start} to {end})"
        ws['A1'].font = title_font

        headers = ["Date", "Customer", "Type", "Description", "Amount (OMR)", "Payment", "Employee"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border

        sales_q = db.query(DailySale).filter(DailySale.date >= start, DailySale.date <= end)
        if current_user.role == "employee":
            sales_q = sales_q.filter(DailySale.employee_id == current_user.id)
        sales = sales_q.order_by(DailySale.date.desc()).all()

        for row, s in enumerate(sales, 4):
            ws.cell(row=row, column=1, value=s.date.isoformat()).border = border
            ws.cell(row=row, column=2, value=s.customer_name or "").border = border
            ws.cell(row=row, column=3, value=s.sale_type).border = border
            ws.cell(row=row, column=4, value=s.description or "").border = border
            ws.cell(row=row, column=5, value=float(s.amount)).border = border
            ws.cell(row=row, column=5).number_format = '#,##0.000'
            ws.cell(row=row, column=6, value=s.payment_method).border = border
            ws.cell(row=row, column=7, value=s.employee.full_name if s.employee else "").border = border

        total_row = len(sales) + 4
        ws.cell(row=total_row, column=4, value="TOTAL").font = Font(bold=True)
        ws.cell(row=total_row, column=5, value=sum(float(s.amount) for s in sales))
        ws.cell(row=total_row, column=5).font = Font(bold=True)
        ws.cell(row=total_row, column=5).number_format = '#,##0.000'

    elif report_type == "expenses":
        ws.title = "Expenses Report"
        ws.merge_cells('A1:G1')
        ws['A1'] = f"Expenses Report ({start} to {end})"
        ws['A1'].font = title_font

        headers = ["Date", "Category", "Description", "Amount (OMR)", "Payment", "Employee", "Notes"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border

        expenses_q = db.query(DailyExpense).filter(DailyExpense.date >= start, DailyExpense.date <= end)
        if current_user.role == "employee":
            expenses_q = expenses_q.filter(DailyExpense.employee_id == current_user.id)
        expenses = expenses_q.order_by(DailyExpense.date.desc()).all()

        for row, e in enumerate(expenses, 4):
            ws.cell(row=row, column=1, value=e.date.isoformat()).border = border
            ws.cell(row=row, column=2, value=e.category).border = border
            ws.cell(row=row, column=3, value=e.description or "").border = border
            ws.cell(row=row, column=4, value=float(e.amount)).border = border
            ws.cell(row=row, column=4).number_format = '#,##0.000'
            ws.cell(row=row, column=5, value=e.payment_method).border = border
            ws.cell(row=row, column=6, value=e.employee.full_name if e.employee else "").border = border
            ws.cell(row=row, column=7, value=e.note or "").border = border

        total_row = len(expenses) + 4
        ws.cell(row=total_row, column=3, value="TOTAL").font = Font(bold=True)
        ws.cell(row=total_row, column=4, value=sum(float(e.amount) for e in expenses))
        ws.cell(row=total_row, column=4).font = Font(bold=True)
        ws.cell(row=total_row, column=4).number_format = '#,##0.000'

    else:  # profit-loss
        ws.title = "Profit & Loss"
        ws.merge_cells('A1:D1')
        ws['A1'] = f"Profit & Loss Report ({start} to {end})"
        ws['A1'].font = title_font

        sales_q = db.query(DailySale).filter(DailySale.date >= start, DailySale.date <= end)
        expenses_q = db.query(DailyExpense).filter(DailyExpense.date >= start, DailyExpense.date <= end)
        if current_user.role == "employee":
            sales_q = sales_q.filter(DailySale.employee_id == current_user.id)
            expenses_q = expenses_q.filter(DailyExpense.employee_id == current_user.id)

        total_sales = to_float(sales_q.with_entities(func.sum(DailySale.amount)).scalar())
        total_expenses = to_float(expenses_q.with_entities(func.sum(DailyExpense.amount)).scalar())

        # Income section
        ws.cell(row=3, column=1, value="INCOME").font = Font(bold=True, size=12)
        headers = ["Type", "Count", "Amount (OMR)"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill

        sales_by_type = (
            sales_q.with_entities(DailySale.sale_type, func.count(), func.sum(DailySale.amount))
            .group_by(DailySale.sale_type).all()
        )
        row = 5
        for t, c, a in sales_by_type:
            ws.cell(row=row, column=1, value=t)
            ws.cell(row=row, column=2, value=c)
            ws.cell(row=row, column=3, value=round(to_float(a), 3))
            ws.cell(row=row, column=3).number_format = '#,##0.000'
            row += 1

        ws.cell(row=row, column=1, value="Total Income").font = Font(bold=True)
        ws.cell(row=row, column=3, value=round(total_sales, 3)).font = Font(bold=True)
        ws.cell(row=row, column=3).number_format = '#,##0.000'
        row += 2

        # Expense section
        ws.cell(row=row, column=1, value="EXPENSES").font = Font(bold=True, size=12)
        row += 1
        for col, h in enumerate(["Category", "Count", "Amount (OMR)"], 1):
            cell = ws.cell(row=row, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
        row += 1

        expenses_by_cat = (
            expenses_q.with_entities(DailyExpense.category, func.count(), func.sum(DailyExpense.amount))
            .group_by(DailyExpense.category)
            .order_by(func.sum(DailyExpense.amount).desc()).all()
        )
        for cat, c, a in expenses_by_cat:
            ws.cell(row=row, column=1, value=cat)
            ws.cell(row=row, column=2, value=c)
            ws.cell(row=row, column=3, value=round(to_float(a), 3))
            ws.cell(row=row, column=3).number_format = '#,##0.000'
            row += 1

        ws.cell(row=row, column=1, value="Total Expenses").font = Font(bold=True)
        ws.cell(row=row, column=3, value=round(total_expenses, 3)).font = Font(bold=True)
        ws.cell(row=row, column=3).number_format = '#,##0.000'
        row += 2

        # Net profit
        profit_fill = PatternFill(start_color="00A09D", end_color="00A09D", fill_type="solid")
        ws.cell(row=row, column=1, value="NET PROFIT").font = Font(bold=True, size=12)
        profit_cell = ws.cell(row=row, column=3, value=round(total_sales - total_expenses, 3))
        profit_cell.font = Font(bold=True, size=12)
        profit_cell.number_format = '#,##0.000'

    # Set column widths
    for col in range(1, 8):
        ws.column_dimensions[chr(64 + col)].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"{report_type}_report_{start}_{end}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
